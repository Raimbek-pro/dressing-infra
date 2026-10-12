"""Builds SIS2_report.docx (14 pt Times New Roman, bold headings) from the repo contents.
Run: python3 -m venv .venv && .venv/bin/pip install python-docx && .venv/bin/python report/build_report.py
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
doc = Document()

for style_name in ("Normal", "Heading 1", "Heading 2", "Title"):
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(14)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if style_name != "Normal":
        style.font.bold = True
doc.styles["Title"].font.size = Pt(18)
doc.styles["Heading 1"].font.size = Pt(16)


def p(text, bold=False, italic=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold, run.italic = bold, italic
    return para


def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def code(text):
    para = doc.add_paragraph()
    run = para.add_run(text.rstrip())
    run.font.name = "Courier New"
    run.font.size = Pt(10)


def table(header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    for row in t.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(11)
                    run.font.name = "Times New Roman"
    doc.add_paragraph()


def screenshot(what):
    p(f"[Screenshot: {what}]", italic=True)


def file(path):
    code((ROOT / path).read_text())


# ---------------------------------------------------------------- title
doc.add_paragraph("Student's Individual Study #2", style="Title")
p("Full name: ____________________   Group: ________   Educational program: ____________________")
p("Repository: https://github.com/<username>/dressing-infra")

doc.add_heading("Name, topic, target and tasks", level=1)
p("Name: SIS #2. Topic: Users and permissions.", bold=False)
p("Target: define a role table and prepare automation of the setup.")
p("System: Linux infrastructure in Google Cloud for Dressing, an iOS app where people "
  "dress garment patterns in fabrics from their own world, share them with friends, gift "
  "collectible pieces and try them on in AR. Two Debian 13 VMs: a backend VM and a "
  "PostgreSQL VM.")
p("Tasks:")
bullets([
    "Declare a role table (roles, permissions, superuser permissions, service accounts).",
    "Create all groups and at least one user per role.",
    "Set file permissions and sudoers permissions for all groups.",
    "Set up SSH key access for the automation user.",
    "Run a service as a systemd unit that starts at every boot.",
])

# ---------------------------------------------------------------- architecture
doc.add_heading("0. Architecture and design decisions", level=1)
doc.add_heading("Why two VMs, and not a third for the frontend", level=2)
p("The frontend is the iOS app itself: it runs on the user's phone and talks to the backend "
  "over HTTPS. Nothing has to be served from a frontend server, so a third VM for it would "
  "cost money and add attack surface without doing work. Images and 3D models (fabric photos, "
  "USDZ files for the AR wardrobe) belong in object storage (a Cloud Storage bucket behind a "
  "CDN), not on a VM disk.")
p("The database runs on its own VM so that it has no public IP at all, its resources (RAM for "
  "the cache, disk I/O) are not shared with the API, and the backend can be rebuilt or scaled "
  "without touching the data. This is the standard two-tier layout:")
code("""iPhone ──HTTPS──▶ dressing-app (public IP, 10.10.0.10)         dressing-db (no public IP, 10.10.0.20)
                  nginx :80/443 → dressing-api :8080 ──TLS:5432──▶ PostgreSQL 17
admin ──SSH key──▶ dressing-app ──ProxyJump──▶ dressing-db        apt updates via Cloud NAT""")
p("A third VM becomes useful later only for a separate job: a bastion host (if SSH should not "
  "reach the app VM directly), a worker that converts designs into 3D models for AR, or a "
  "PostgreSQL streaming replica for failover and backups.")

doc.add_heading("Why PostgreSQL", level=2)
bullets([
    "The data is relational: users own designs, and later follow each other and exchange "
    "gifts. Foreign keys and transactions keep that consistent (a gift moves from one owner "
    "to another in one step, never duplicated or lost).",
    "JSONB stores the editable canvas (garments, pieces, fabric per piece) without a rigid "
    "schema, while everything around it stays relational and indexed.",
    "Role-based access is built in and maps onto the Linux role model (NOLOGIN group roles, "
    "LOGIN roles that inherit them), so least privilege applies inside the database too.",
    "It is open source, packaged in Debian, and has a managed version in Google Cloud "
    "(Cloud SQL for PostgreSQL), so the same schema can be moved there later without rewriting.",
])

doc.add_heading("Network and Google Cloud roles", level=2)
p("Own VPC (10.10.0.0/24) instead of the default network, so only these firewall rules exist:")
table(["Rule", "Port", "From", "To"], [
    ["dressing-allow-web", "80, 443", "internet", "dressing-app"],
    ["dressing-allow-ssh-admin", "22", "admin's IP only", "dressing-app"],
    ["dressing-allow-ssh-iap", "22", "Google IAP range", "both VMs"],
    ["dressing-allow-ssh-jump", "22", "dressing-app", "dressing-db"],
    ["dressing-allow-pg-from-app", "5432", "dressing-app", "dressing-db"],
])
p("On the instance creation page, 'Allow HTTP/HTTPS traffic' is ticked only for dressing-app. "
  "For dressing-db it is unticked and the external IP is set to None.")
p("Permissions exist on two layers. Google Cloud IAM decides who may create, stop or SSH "
  "into VMs and read cloud logs. Linux users, groups and sudoers decide what someone can do "
  "inside a VM. Each VM also runs under its own cloud service account that can only write "
  "logs and metrics, instead of the default Compute Engine account, which has Editor on the "
  "whole project.")
table(["Google Cloud principal", "Roles", "Why"], [
    ["Administrator (person)", "compute.admin, iam.serviceAccountAdmin, resourcemanager.projectIamAdmin",
     "Manages VMs, network, IAM. Not Owner: Owner also controls billing and project deletion."],
    ["DevOps (person)", "compute.instanceAdmin.v1, iam.serviceAccountUser, iap.tunnelResourceAccessor",
     "Starts/stops VMs and SSHs through IAP; cannot change firewall or IAM."],
    ["Auditor (person)", "viewer, logging.viewer, iam.securityReviewer", "Read-only, including logs and IAM policies."],
    ["sa-dressing-app", "logging.logWriter, monitoring.metricWriter", "Identity of the app VM."],
    ["sa-dressing-db", "logging.logWriter, monitoring.metricWriter", "Identity of the DB VM."],
])
p("Script: gcp/01_create_infra.sh. The people roles above are shown as the target setup for a team.")
screenshot("VM instances list showing dressing-app (with external IP) and dressing-db (no external IP)")
screenshot("Firewall rules of dressing-vpc")

# ---------------------------------------------------------------- task 1
doc.add_heading("1. Role table", level=1)
p("Linux roles. Service accounts have no password and the shell /usr/sbin/nologin. "
  "They exist so each program runs with only its own permissions.")
table(["Role (group / user)", "Description", "Permissions", "Superuser permissions (sudo)"], [
    ["sysadmins / sysadmin1", "System administrator", "Own home; member of adm (reads /var/log)",
     "ALL, with own password every time (5-minute cache)"],
    ["devops / devops1", "Operates the backend and web server", "Own home; member of developers",
     "systemctl start|stop|restart|status dressing-api; systemctl reload nginx; nginx -t"],
    ["developers / developer1", "Backend developer", "Own home only", "systemctl status dressing-api"],
    ["dba / dba1", "Database administrator", "DB role dba1 (schema owner, peer auth)",
     "psql, pg_dump as user postgres; restart/status postgresql"],
    ["auditors / auditor1", "Security auditor, read-only", "systemd-journal + adm groups (reads all logs)",
     "NOPASSWD, NOEXEC: cat /var/log/sudo.log, sudo -l -U <user>"],
    ["automation / automation_bot", "CI/CD and Ansible; SSH key only, no password",
     "Member of devops", "NOPASSWD: systemctl start|stop|restart|status dressing-api"],
    ["dressing (service)", "Runs the backend API", "Read /opt/dressing, /etc/dressing; write only /var/lib/dressing",
     "None; also blocked by systemd sandboxing (NoNewPrivileges, ProtectSystem=strict)"],
    ["postgres (service)", "Runs PostgreSQL (created by the package)", "Owns /var/lib/postgresql (0700) and /etc/postgresql",
     "None. DB superuser inside PostgreSQL, reachable only through the local socket (peer auth)"],
    ["www-data (service)", "Runs nginx workers (created by the package)", "Read site config; no write access to the app",
     "None"],
    ["root", "Kernel-level super user", "Everything", "Direct login disabled (PermitRootLogin no)"],
])
p("Database roles (inside PostgreSQL, same group/user idea):")
table(["Role", "Login", "Privileges"], [
    ["dressing_owner", "no", "Owns database dressing and schema app; runs migrations (DDL)"],
    ["dressing_readwrite", "no", "SELECT, INSERT, UPDATE, DELETE on app tables; sequences"],
    ["dressing_readonly", "no", "SELECT on app tables (analytics, auditors)"],
    ["dressing_app", "yes, SCRAM password, max 20 connections", "Member of dressing_readwrite; only from 10.10.0.10 over TLS"],
    ["dba1", "yes, peer (OS user dba1 only)", "Member of dressing_owner"],
    ["postgres", "yes, peer, local socket only", "Superuser"],
])

# ---------------------------------------------------------------- task 2
doc.add_heading("2. Groups and users", level=1)
p("Script linux/10_users_groups.sh (run on both VMs). It creates the six groups, the dressing "
  "service account (system user, nologin) and one user per role. Human accounts start locked "
  "(passwd -l) until an administrator sets a password; automation_bot never gets one. Home "
  "directories are set to 0750 so users cannot read each other's files.")
file("linux/10_users_groups.sh")
screenshot("output of the script: getent group ... and id <user> for each user")
screenshot("tail /etc/passwd and /etc/group")

# ---------------------------------------------------------------- task 3
doc.add_heading("3. File and sudoers permissions", level=1)
table(["Path", "Owner:group", "Mode", "Reason"], [
    ["/opt/dressing", "root:dressing", "0750", "Code: the service can run it but not modify itself"],
    ["/etc/dressing/dressing.env", "root:dressing", "0640", "DATABASE_URL secret: only root writes, only the service reads"],
    ["/var/lib/dressing", "dressing:dressing", "0750", "The only directory the service may write"],
    ["/etc/sudoers.d/10-dressing", "root:root", "0440", "Sudo policy, validated with visudo -c before install"],
    ["/var/log/sudo.log", "root:root", "0600", "Every sudo call; readable by auditors only through sudo"],
])
p("The sudo policy is placed in /etc/sudoers.d instead of editing /etc/sudoers, and is "
  "checked with visudo before it is installed. Command aliases keep it readable. Every "
  "group gets only the exact commands its job needs; DevOps and the bot never get a root "
  "shell, and the bot (no password) can run only the app's service commands.")
file("linux/sudoers.d/10-dressing")
file("linux/20_permissions.sh")
screenshot("ls -ld of the directories above")
screenshot("sudo -l -U for each user (printed at the end of 20_permissions.sh)")
screenshot("a denied command, e.g. developer1 running sudo systemctl restart dressing-api")

# ---------------------------------------------------------------- task 4
doc.add_heading("4. SSH with a key pair as the automation user", level=1)
p("The key pair is generated on the administrator's laptop (ed25519); only the public key "
  "goes to the servers. sshd gets a drop-in config that disables root login and passwords, "
  "and only lets the role groups log in. The config is validated with sshd -t before "
  "reloading, so a typo cannot lock everyone out. The DB VM has no public IP; it is reached "
  "through the app VM with ProxyJump.")
code("ssh-keygen -t ed25519 -C automation_bot -f ~/.ssh/dressing_automation_bot")
file("linux/30_ssh.sh")
code("""Host dressing-app
    HostName <app external IP>
    User automation_bot
    IdentityFile ~/.ssh/dressing_automation_bot
Host dressing-db
    HostName 10.10.0.20
    User automation_bot
    IdentityFile ~/.ssh/dressing_automation_bot
    ProxyJump dressing-app""")
screenshot("ssh dressing-app and ssh dressing-db as automation_bot without a password")
screenshot("password login refused: ssh -o PubkeyAuthentication=no automation_bot@<ip>")

# ---------------------------------------------------------------- task 5
doc.add_heading("5. PostgreSQL configuration", level=1)
p("PostgreSQL 17 from Debian. Our settings go into conf.d so package upgrades do not "
  "conflict with them. listen_addresses includes only localhost and the private IP. "
  "Passwords use SCRAM-SHA-256, connections use TLS, and memory is sized for a 1 GB VM.")
file("db/postgresql.d/10-dressing.conf")
p("pg_hba.conf decides who may connect from where; the first matching line wins and "
  "everything else is rejected:")
file("db/pg_hba.conf")
file("db/roles.sql")
screenshot("\\du in psql and ss -ltnp | grep 5432")
screenshot("psql as dressing_app: SELECT works, CREATE TABLE gives 'permission denied for schema app'")

doc.add_heading("6. The backend as a systemd unit", level=1)
p("The service is a small Python API (app/dressing_api.py) that answers /health by running "
  "SELECT 1 on the database VM. It listens on 127.0.0.1:8080; nginx (www-data) is the only "
  "process exposed to the internet. The unit runs as the dressing user with systemd "
  "sandboxing and is enabled, so it starts at every boot.")
file("app/dressing-api.service")
code("sudo systemctl enable --now dressing-api\nsudo reboot\ncurl http://<app external IP>/health   # {\"status\": \"ok\", \"db\": \"ok\"}")
screenshot("systemctl status dressing-api after reboot (enabled; active (running))")
screenshot("curl /health from the laptop")

# ---------------------------------------------------------------- future
doc.add_heading("7. How the infrastructure grows with the app", level=1)
p("The schema (db/schema.sql) starts with users and their designs. Planned next:")
bullets([
    "Community tab: following friends and a feed of public designs.",
    "Gifts as collectible items that move between users, recorded as transactions.",
    "AR wardrobe: 3D versions of designs stored in Cloud Storage and fitted to the body on "
    "the phone. Generating the 3D models is the first job that would justify a third machine.",
])
p("Scaling path, without redesigning: Cloud Storage plus Cloud CDN for photos and models; a "
  "PostgreSQL streaming replica (wal_level = replica is already set) or a move to Cloud SQL; "
  "several app VMs behind a load balancer (the API is stateless); secrets moved from the env "
  "file to Secret Manager; OS Login so Linux accounts follow Google Cloud IAM automatically.")

# ---------------------------------------------------------------- end
doc.add_heading("Conclusions", level=1)
p("The system separates who manages the cloud (IAM) from who does what inside each machine "
  "(Linux groups and sudoers) and inside the database (PostgreSQL roles). Every program runs "
  "as its own service account, no human or bot gets more than its job needs, root login and "
  "passwords over SSH are disabled, the database is unreachable from the internet, and the "
  "whole setup is reproducible from scripts in the repository.")

out = ROOT / "report" / "SIS2_report.docx"
doc.save(out)
print("saved", out)
